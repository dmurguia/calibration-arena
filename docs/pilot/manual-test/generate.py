"""Regenerate the synthetic manual-test files and print the answer-key figures.

Run from the repository root with the repo virtualenv (openpyxl is a backend dependency):

    .venv/bin/python docs/pilot/manual-test/generate.py

Edit a scenario's data below, rerun, and copy the printed figures into README.md.
All companies, people and amounts are fictional.
"""
import csv
import json
import sys
from datetime import date, timedelta
from decimal import Decimal as D
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).parent


# ---------------------------------------------------------------- Bank reconciliation, September 2026
# Prior-month reconciliation (deliberately not supplied): at August 31 checks 1088 ($1,240.00) and
# 1091 ($615.40) were outstanding and a $3,180.00 deposit was in transit.
BANK_OPEN, BOOK_OPEN = D('48210.55'), D('49535.15')

# (date, description, signed amount, on bank statement, in cash ledger)
CASH = [
    ('2026-09-01', 'Deposit', '3180.00', True, False),
    ('2026-09-02', 'Square settlement', '3412.18', True, True),
    ('2026-09-03', 'Check 1088', '-1240.00', True, False),
    ('2026-09-04', 'Check 1101 - Pacific Bean Importers', '-8450.00', True, True),
    ('2026-09-05', 'ACH Gusto payroll', '-11872.36', True, True),
    ('2026-09-05', 'Square settlement', '2987.05', True, True),
    ('2026-09-08', 'Wire in - Northside Grocers inv 2231', '6240.00', True, True),
    ('2026-09-09', 'Check 1102 - PG&E', '-1128.44', True, True),
    ('2026-09-10', 'Check 1103 - Harborview Properties rent', '-6500.00', True, True),
    ('2026-09-12', 'Square settlement', '3105.77', True, True),
    ('2026-09-12', 'Deposit - Bayside Cafe check 4471 inv 2219', '1150.00', True, True),
    ('2026-09-15', 'Check 1104 - Coastal Packaging', '-2214.90', True, True),
    ('2026-09-16', 'Returned item - Bayside Cafe chk 4471 NSF', '-1150.00', True, False),
    ('2026-09-16', 'Returned item fee', '-35.00', True, False),
    ('2026-09-17', 'Check 1105 - Valley Dairy Supply', '-1254.00', True, True),
    ('2026-09-19', 'ACH Gusto payroll', '-11940.12', True, True),
    ('2026-09-19', 'Square settlement', '2860.42', True, True),
    ('2026-09-22', 'ACH debit KDZ HOLDINGS LLC WEB PMT', '-1980.00', True, False),
    ('2026-09-23', 'Deposit - Northside Grocers inv 2236', '4815.50', True, True),
    ('2026-09-24', 'Check 1106 - Pacific Bean Importers', '-7925.00', True, True),
    ('2026-09-25', 'Check 1107 - Comcast Business', '-289.99', True, True),
    ('2026-09-26', 'Square settlement', '3640.11', True, True),
    ('2026-09-26', 'Check 1108 - CDTFA sales tax', '-2746.00', True, True),
    ('2026-09-28', 'Check 1112 - Mission Paper Goods', '-2318.75', False, True),
    ('2026-09-29', 'Check 1109 - Bay Area Roaster Repair', '-735.00', True, True),
    ('2026-09-29', 'Check 1111 - Sonoma Sweets', '-418.20', True, True),
    ('2026-09-29', 'Square settlement', '3018.64', True, True),
    ('2026-09-30', 'Check 1114 - Delta Electric', '-960.00', False, True),
    ('2026-09-30', 'Deposit - Northside Grocers inv 2240', '4275.60', False, True),
    ('2026-09-30', 'Interest earned', '12.38', True, False),
    ('2026-09-30', 'Monthly service fee', '-45.00', True, False),
]
LEDGER_AMOUNT = {'Check 1105 - Valley Dairy Supply': D('-1524.00')}  # transposed when booked
LEDGER_DUPLICATE = ('2026-09-19', 'Square settlement')                # posted twice
LEDGER_VOID = ('2026-09-27', 'Check 1110 - VOID', D('0.00'))          # 1113 is missing from both files


def bank_reconciliation():
    bank = [(d, desc, D(a)) for d, desc, a, on_bank, _ in CASH if on_bank]
    ledger = []
    for d, desc, a, _, in_ledger in CASH:
        if in_ledger:
            ledger.append((d, desc, LEDGER_AMOUNT.get(desc, D(a))))
            if (d, desc) == LEDGER_DUPLICATE:
                ledger.append(ledger[-1])
    ledger.insert(next(i for i, r in enumerate(ledger) if r[1].startswith('Check 1109')), LEDGER_VOID)

    bank_close = BANK_OPEN + sum(a for *_, a in bank)
    book_close = BOOK_OPEN + sum(a for *_, a in ledger)
    adjusted_bank = bank_close + D('4275.60') - D('2318.75') - D('960.00') - D('615.40')
    adjusted_book = (book_close - D('45.00') + D('12.38') - D('1150.00') - D('35.00')
                     + D('270.00') - D('2860.42') - D('1980.00'))
    assert adjusted_bank == adjusted_book, 'the reconciliation no longer ties; update the answer key'

    with open(OUT / 'bank-statement.csv', 'w', newline='') as f:
        w = csv.writer(f)
        w.writerow(['Date', 'Description', 'Debit', 'Credit', 'Balance'])
        balance = BANK_OPEN
        w.writerow(['2026-09-01', 'Opening balance', '', '', f'{balance:.2f}'])
        for d, desc, a in bank:
            balance += a
            w.writerow([d, desc.upper(), f'{-a:.2f}' if a < 0 else '', f'{a:.2f}' if a > 0 else '', f'{balance:.2f}'])
        w.writerow(['2026-09-30', 'Closing balance', '', '', f'{balance:.2f}'])

    with open(OUT / 'cash-ledger.csv', 'w', newline='') as f:
        w = csv.writer(f)
        w.writerow(['Date', 'GL Account', 'Reference', 'Description', 'Debit', 'Credit', 'Running Balance'])
        balance = BOOK_OPEN
        w.writerow(['2026-08-31', '1010 Cash - Operating', 'BAL', 'Balance forward', '', '', f'{balance:.2f}'])
        for d, desc, a in ledger:
            balance += a
            reference = (desc.split(' - ')[0] if desc.startswith('Check') else 'SQ' if 'Square' in desc
                         else 'DEP' if 'Deposit' in desc or 'Wire' in desc else 'ACH')
            w.writerow([d, '1010 Cash - Operating', reference, desc,
                        f'{a:.2f}' if a > 0 else '', f'{-a:.2f}' if a < 0 else '', f'{balance:.2f}'])

    credits, debits = sum(a for *_, a in bank if a > 0), -sum(a for *_, a in bank if a < 0)
    lines = ['PACIFIC COAST COMMUNITY BANK', 'Business Checking Statement', '',
             'Harbor & Pine Coffee Roasters LLC          Account: ****4417',
             'Statement period: September 1 - September 30, 2026', '',
             f'Beginning balance 09/01/2026   {BANK_OPEN:>12,.2f}',
             f'Deposits and credits           {credits:>12,.2f}',
             f'Withdrawals and debits         {debits:>12,.2f}',
             f'Ending balance 09/30/2026      {bank_close:>12,.2f}', '',
             'DATE   DESCRIPTION                                       AMOUNT       BALANCE',
             '-' * 80]
    balance = BANK_OPEN
    for d, desc, a in bank:
        balance += a
        lines.append(f'{d[5:7]}/{d[8:]}  {desc.upper()[:46]:<46} {a:>12,.2f}  {balance:>12,.2f}')
    lines += ['-' * 80, '', 'Checks paid: 1088 1101 1102 1103 1104 1105 1106 1107 1108 1109 1111',
              '* Indicates a break in check sequence: 1088 *1101 ... 1109 *1111', '',
              'Questions? Call 1-800-555-0142. This is a synthetic statement for testing.']
    write_text_pdf(OUT / 'bank-statement.pdf', lines)
    return {'bank_close': bank_close, 'book_close': book_close, 'adjusted_balance': adjusted_bank,
            'bank_credits': credits, 'bank_debits': debits}


def write_text_pdf(path, lines):
    """One Courier page with extractable text; hand-built so no PDF library is needed."""
    def escape(s):
        return s.replace('\\', '\\\\').replace('(', '\\(').replace(')', '\\)')
    stream = 'BT /F1 8.5 Tf 11 TL 36 760 Td\n' + '\n'.join(f'({escape(l)}) Tj T*' for l in lines) + '\nET'
    objects = ['<< /Type /Catalog /Pages 2 0 R >>', '<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
               '<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>',
               '<< /Type /Font /Subtype /Type1 /BaseFont /Courier >>',
               f'<< /Length {len(stream.encode())} >>\nstream\n{stream}\nendstream']
    pdf, offsets = b'%PDF-1.4\n', []
    for number, body in enumerate(objects, 1):
        offsets.append(len(pdf))
        pdf += f'{number} 0 obj\n{body}\nendobj\n'.encode()
    xref = len(pdf)
    pdf += f'xref\n0 {len(objects) + 1}\n0000000000 65535 f \n'.encode()
    pdf += b''.join(f'{o:010d} 00000 n \n'.encode() for o in offsets)
    pdf += f'trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n'.encode()
    path.write_bytes(pdf)


# ---------------------------------------------------------------- AR reporting, aging at September 30, 2026
REPORTING_DATE = date(2026, 9, 30)
GL_DIFFERENCE = D('475.00')  # GL control exceeds the subledger by an unexplained amount
TERMS = {'Northside Grocers Co-op': 30, 'Bayside Cafe': 15, 'Mission District Bagels': 30, 'Sunset Hotel Group': 45,
         'Golden Gate Coworking': 30, 'Marin Farm Market': 30, 'Telegraph Ave Books & Brew': 30}
# (document, customer, invoice date, original, applied, due date override: None = terms, 'blank' = empty, note)
AR = [
    ('INV-2188', 'Northside Grocers Co-op', '2026-06-12', '5120.00', '0.00', None, ''),
    ('INV-2214', 'Northside Grocers Co-op', '2026-08-20', '4980.00', '0.00', None, ''),
    ('INV-2236', 'Northside Grocers Co-op', '2026-09-02', '4815.50', '4815.50', None, 'Paid 09/23'),
    ('INV-2240', 'Northside Grocers Co-op', '2026-09-15', '4275.60', '0.00', None, ''),
    ('INV-2246', 'Northside Grocers Co-op', '2026-09-28', '5310.25', '0.00', None, ''),
    ('INV-2172', 'Bayside Cafe', '2026-05-28', '1460.00', '0.00', None, 'Customer on credit hold'),
    ('INV-2207', 'Bayside Cafe', '2026-07-30', '1385.40', '400.00', None, 'Partial payment 08/25'),
    ('INV-2219', 'Bayside Cafe', '2026-08-26', '1150.00', '1150.00', None, 'Paid 09/12 chk 4471'),
    ('INV-2241', 'Bayside Cafe', '2026-09-17', '1290.10', '0.00', None, ''),
    ('INV-2161', 'Mission District Bagels', '2026-05-15', '2240.00', '0.00', None, 'Disputed - short shipment claim 6/02'),
    ('INV-2222', 'Mission District Bagels', '2026-08-28', '2615.80', '0.00', None, ''),
    ('INV-2243', 'Mission District Bagels', '2026-09-22', '2480.00', '0.00', None, ''),
    ('UNAPPLIED', 'Mission District Bagels', '2026-09-24', '-1200.00', '0.00', 'blank', 'ACH received, no remittance'),
    ('INV-2195', 'Sunset Hotel Group', '2026-07-01', '9850.00', '0.00', None, ''),
    ('INV-2215', 'Sunset Hotel Group', '2026-08-21', '10420.00', '0.00', None, ''),
    ('INV-2238', 'Sunset Hotel Group', '2026-09-10', '11275.00', '0.00', None, ''),
    ('CM-0412', 'Sunset Hotel Group', '2026-09-18', '-350.00', '0.00', 'blank', 'Credit for damaged bags, apply to INV-2215'),
    ('INV-2201', 'Golden Gate Coworking', '2026-07-18', '1875.00', '0.00', '2026-09-16', ''),  # disagrees with Net 30
    ('INV-2229', 'Golden Gate Coworking', '2026-09-01', '1940.00', '0.00', None, ''),
    ('INV-2247', 'Golden Gate Coworking', '2026-09-29', '1990.00', '0.00', None, ''),
    ('INV-2129', 'Marin Farm Market', '2026-04-03', '3180.00', '0.00', None, 'Sent to collections 8/15'),
    ('INV-2210', 'Marin Farm Market', '2026-08-12', '2765.00', '0.00', 'blank', ''),
    ('INV-2244', 'Marin Farm Market', '2026-09-24', '2890.00', '0.00', None, ''),
    ('INV-2224', 'Telegraph Ave Books & Brew', '2026-08-29', '845.00', '0.00', None, ''),
    ('INV-2245', 'Telegraph Ave Books & Brew', '2026-09-26', '910.00', '0.00', None, ''),
]
BUCKETS = ['Current', '1-30', '31-60', '61-90', 'Over 90']


def ar_reporting():
    by_customer = {c: dict.fromkeys(BUCKETS, D(0)) for c in TERMS}
    rows = []
    for document, customer, invoice_date, original, applied, due, note in AR:
        invoiced = date.fromisoformat(invoice_date)
        by_terms = invoiced + timedelta(days=TERMS[customer])
        recorded_due = None if due == 'blank' else date.fromisoformat(due) if due else by_terms
        open_balance = D(original) - D(applied)
        rows.append((document, customer, invoiced, f'Net {TERMS[customer]}', recorded_due,
                     D(original), D(applied), open_balance, note))
        if open_balance:
            # The key ages invoices by their terms and shows unapplied credits as current.
            days = (REPORTING_DATE - (by_terms if document.startswith('INV') else REPORTING_DATE)).days
            bucket = ('Current' if days <= 0 else '1-30' if days <= 30 else '31-60' if days <= 60
                      else '61-90' if days <= 90 else 'Over 90')
            by_customer[customer][bucket] += open_balance
    total = sum(r[7] for r in rows)

    book = Workbook()
    sheet = book.active
    sheet.title = 'AR Open Items'
    sheet.append(['Harbor & Pine Coffee Roasters LLC - Accounts Receivable Open Items'])
    sheet['A1'].font = Font(bold=True)
    sheet.append(['Reporting date', REPORTING_DATE])
    sheet['B2'].number_format = 'yyyy-mm-dd'
    sheet.append([])
    sheet.append(['Document', 'Customer', 'Invoice Date', 'Terms', 'Due Date', 'Original Amount',
                  'Payments/Credits Applied', 'Open Balance', 'Notes'])
    for cell in sheet[4]:
        cell.font = Font(bold=True)
    for document, customer, invoiced, terms, due, original, applied, open_balance, note in rows:
        sheet.append([document, customer, invoiced, terms, due, float(original), float(applied), float(open_balance), note])
        for column in 'CE':
            sheet[f'{column}{sheet.max_row}'].number_format = 'yyyy-mm-dd'
        for column in 'FGH':
            sheet[f'{column}{sheet.max_row}'].number_format = '#,##0.00'
    sheet.append([])
    sheet.append(['Total', None, None, None, None, None, None, f'=SUM(H5:H{sheet.max_row - 1})'])
    sheet[f'H{sheet.max_row}'].font = Font(bold=True)
    for column, width in zip('ABCDEFGHI', [12, 28, 12, 8, 12, 15, 22, 14, 44]):
        sheet.column_dimensions[column].width = width
    gl = book.create_sheet('GL Tie-out')
    gl.append(['Account', 'Description', 'Balance 2026-09-30'])
    gl.append(['1200', 'Accounts Receivable - Trade', float(total + GL_DIFFERENCE)])
    gl.append(['1210', 'Allowance for Doubtful Accounts', -4200.00])
    gl.append([])
    gl.append(['Trailing 3-month net credit sales (Jul-Sep 2026)', None, 118460.00])
    book.save(OUT / 'ar-ledger.xlsx')
    return {'total_open': total, 'gl_balance': total + GL_DIFFERENCE,
            'bucket_totals': {b: sum(by_customer[c][b] for c in TERMS) for b in BUCKETS},
            'by_customer': by_customer}


# ---------------------------------------------------------------- Invoicing, October 2026
BILLING_TERMS = '''LEDGERLINE ADVISORY LLC - STATEMENT OF WORK #3 (EXCERPT)
Client: Sunset Hotel Group LLC, 1450 Great Highway, San Francisco, CA 94122
Client AP contact: ap@sunsethotelgroup.example
Client purchase order: PO-SHG-7781 (must appear on every invoice; invoices without it will be rejected)
SOW term: July 1, 2026 - December 31, 2026. Invoiced monthly in arrears for the calendar month.

1. Fees
1.1 Time and materials rates (per hour, billed in 0.25-hour increments):
    Engagement Director   $240.00
    Senior Associate      $165.00
    Associate             $110.00
1.2 ERP support retainer: $3,000.00 per month, fixed. The retainer includes up to 10.0 hours of
    Associate time per month. Associate hours above 10.0 in a month are billed at the Associate rate.
    Retainer hours do not roll over.
1.3 Monthly fee cap: time-and-materials fees (excluding the retainer, expenses and software) shall
    not exceed $18,000.00 in any calendar month without a written change order signed by both parties.
1.4 Rates are fixed for the SOW term. Any rate other than those in 1.1 requires a written change order.

2. Expenses
2.1 Travel is reimbursable at actual cost with no markup, only if pre-approved in writing by the
    client project sponsor (R. Okafor). Reimbursable travel is capped at $750.00 per month.
2.2 Mileage is reimbursable at $0.70 per mile and counts toward the travel cap.

3. Software resale
3.1 Planning software seats resold to the client are billed at $95.00 per seat per month.
3.2 Professional services and the retainer are not subject to sales tax (client instruction, 7/1/2026).
3.3 Software seats are subject to sales tax at 8.625% unless a valid resale certificate from the client
    is on file before the invoice date.

4. Payment terms
4.1 Net 30 from invoice date. 2% discount if paid in full within 10 days (2/10, net 30).
4.2 Late balances accrue 1.0% per month.

CHANGE ORDER CO-2 (signed by both parties October 20, 2026)
Section 1.3 is amended: the monthly fee cap is increased to $22,000.00, effective October 20, 2026.
All other terms unchanged.
'''
RATES = {'Engagement Director': D(240), 'Senior Associate': D(165), 'Associate': D(110)}
RETAINER, RETAINER_HOURS, TRAVEL_CAP, SEAT_PRICE, SEATS, TAX_RATE = D(3000), 10, D(750), D(95), 6, D('0.08625')
# (date, staff, role or type, hours, description, client approved, billable, notes)
TIME = [
    ('2026-09-30', 'M. Alvarez', 'Senior Associate', 3.0, 'September close support - accrual review', 'Y', 'Y', 'Logged late'),
    ('2026-10-01', 'J. Kim', 'Engagement Director', 2.0, 'Monthly steering meeting', 'Y', 'Y', ''),
    ('2026-10-02', 'M. Alvarez', 'Senior Associate', 6.5, 'Revenue cut-off testing', 'Y', 'Y', ''),
    ('2026-10-03', 'T. Nguyen', 'Associate', 4.0, 'ERP support - vendor master cleanup', 'Y', 'Y', ''),
    ('2026-10-06', 'M. Alvarez', 'Senior Associate', 7.0, 'Lease schedule rebuild', 'Y', 'Y', ''),
    ('2026-10-07', 'T. Nguyen', 'Associate', 3.5, 'ERP support - user access tickets', 'Y', 'Y', ''),
    ('2026-10-08', 'J. Kim', 'Engagement Director', 3.0, 'Audit committee prep', 'Y', 'Y', ''),
    ('2026-10-09', 'M. Alvarez', 'Senior Associate', 8.0, 'Lease schedule rebuild', 'Y', 'Y', ''),
    ('2026-10-10', 'T. Nguyen', 'Associate', 2.0, 'Internal training - new ERP module', 'Y', 'N', 'Ledgerline internal'),
    ('2026-10-13', 'M. Alvarez', 'Senior Associate', 7.5, 'ASC 842 memo draft', 'Y', 'Y', ''),
    ('2026-10-14', 'T. Nguyen', 'Associate', 4.5, 'ERP support - month-end batch failures', 'Y', 'Y', ''),
    ('2026-10-15', 'M. Alvarez', 'Senior Associate', 7.5, 'ASC 842 memo draft', 'Y', 'Y', ''),
    ('2026-10-15', 'M. Alvarez', 'Senior Associate', 7.5, 'ASC 842 memo draft', 'Y', 'Y', ''),
    ('2026-10-16', 'J. Kim', 'Engagement Director', 4.0, 'Memo review and client call', 'Y', 'Y', ''),
    ('2026-10-20', 'T. Nguyen', 'Associate', 6.0, 'Fixed asset subledger reconciliation', 'Y', 'Y', 'Bill at Senior rate per J. Kim'),
    ('2026-10-21', 'M. Alvarez', 'Senior Associate', 8.0, 'Fixed asset subledger reconciliation', 'Y', 'Y', ''),
    ('2026-10-22', 'M. Alvarez', 'Senior Associate', 8.0, 'Impairment indicators analysis', 'Y', 'Y', ''),
    ('2026-10-23', 'J. Kim', 'Engagement Director', 3.5, 'Audit committee presentation', 'Y', 'Y', ''),
    ('2026-10-26', 'M. Alvarez', 'Senior Associate', 8.0, 'Audit PBC support', 'Y', 'Y', ''),
    ('2026-10-26', 'J. Kim', 'Engagement Director', 3.0, 'Planning meeting with client auditors', 'Y', 'Y', ''),
    ('2026-10-27', 'M. Alvarez', 'Senior Associate', 6.0, 'Audit PBC support', 'Y', 'Y', ''),
    ('2026-10-28', 'M. Alvarez', 'Senior Associate', 7.0, 'Impairment memo draft', 'Y', 'Y', ''),
    ('2026-10-28', 'T. Nguyen', 'Associate', 3.0, 'ERP support - report writer fixes', 'N', 'Y', 'Awaiting client approval'),
    ('2026-10-29', 'M. Alvarez', 'Senior Associate', 5.5, 'Audit PBC support', 'Y', 'Y', ''),
    ('2026-10-29', 'J. Kim', 'Engagement Director', 2.5, 'Impairment memo review', 'Y', 'Y', ''),
    ('2026-10-30', 'J. Kim', 'Engagement Director', 2.0, 'Monthly steering meeting', 'Y', 'Y', ''),
]
EXPENSES = [
    ('2026-10-08', 'J. Kim', 'Expense', 0, 'Mileage SF office to client HQ, 2 x 38 mi round trips', 'Y', 'Y', '76 miles; pre-approved R. Okafor 10/01'),
    ('2026-10-22', 'M. Alvarez', 'Expense', 0, 'Hotel - Sacramento property site visit (1 night)', 'Y', 'Y', '$412.80; pre-approved R. Okafor 10/15'),
    ('2026-10-22', 'M. Alvarez', 'Expense', 0, 'Airfare SFO-SMF round trip', 'Y', 'Y', '$318.40; pre-approved R. Okafor 10/15'),
    ('2026-10-23', 'J. Kim', 'Expense', 0, 'Client dinner - audit committee', 'Y', 'Y', '$186.25; no pre-approval on file'),
    ('2026-10-31', 'Ledgerline', 'Software', 0, 'Adaptive planning seats - October', 'Y', 'Y', '6 seats; client says resale certificate is coming, not received as of 10/31'),
]
APPROVED_TRAVEL = D('76') * D('0.70') + D('412.80') + D('318.40')


def invoicing():
    (OUT / 'billing-terms.txt').write_text(BILLING_TERMS)
    with open(OUT / 'billable-items.csv', 'w', newline='') as f:
        w = csv.writer(f)
        w.writerow(['Date', 'Staff', 'Role / Type', 'Hours', 'Description', 'Client Approved', 'Billable', 'Notes'])
        for r in TIME + EXPENSES:
            w.writerow([r[0], r[1], r[2], f'{r[3]:.2f}' if r[3] else '', *r[4:]])

    # The key excludes September time, unapproved time, non-billable time and exact duplicates,
    # and bills the "Senior rate" note at the contractual Associate rate.
    seen, kept, duplicates = set(), [], []
    for r in TIME:
        if r[0] >= '2026-10-01' and r[5] == 'Y' and r[6] == 'Y':
            (duplicates if r[:5] in seen else kept).append(r)
            seen.add(r[:5])

    def fees(entries):
        hours = {role: sum(D(str(r[3])) for r in entries if r[2] == role) for role in RATES}
        hours['Associate'] = max(D(0), hours['Associate'] - RETAINER_HOURS)
        return sum(hours[role] * RATES[role] for role in RATES)

    time_and_materials = fees(kept)
    software = SEATS * SEAT_PRICE
    tax = (software * TAX_RATE).quantize(D('0.01'))
    travel = min(APPROVED_TRAVEL, TRAVEL_CAP)
    return {'hours': {role: sum(D(str(r[3])) for r in kept if r[2] == role) for role in RATES},
            'time_and_materials': time_and_materials,
            'time_and_materials_before_oct_20': fees([r for r in kept if r[0] < '2026-10-20']),
            'duplicate_value': sum(D(str(r[3])) * RATES[r[2]] for r in duplicates),
            'travel_incurred': APPROVED_TRAVEL, 'travel_billed': travel, 'software': software, 'tax': tax,
            'total': RETAINER + time_and_materials + travel + software + tax}


if __name__ == '__main__':
    OUT.mkdir(parents=True, exist_ok=True)
    key = {'bank_reconciliation': bank_reconciliation(), 'ar_reporting': ar_reporting(), 'invoicing': invoicing()}
    print(json.dumps(key, indent=1, default=str))
